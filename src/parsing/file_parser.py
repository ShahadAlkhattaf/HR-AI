"""Extract native text from PDF and DOCX resumes and flag pages needing OCR."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional


class UnsupportedFileTypeError(ValueError):
    pass


@dataclass
class ParsedPage:
    page_number: int
    text: str
    char_count: int = field(init=False)

    def __post_init__(self):
        self.char_count = len(self.text.strip())


@dataclass
class ParseResult:
    source_path: str
    file_type: str  # "pdf" | "docx"
    pages: List[ParsedPage]
    extraction_method: str  # "native" — set to "ocr" by the OCR fallback stage
    needs_ocr: bool

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages)

    @property
    def total_chars(self) -> int:
        return sum(p.char_count for p in self.pages)


# Sparse native text can indicate a scan or broken font extraction; try OCR.
MIN_CHARS_PER_PAGE_THRESHOLD = 20


def parse_pdf(path: str) -> ParseResult:
    import fitz  # PyMuPDF

    doc = fitz.open(path)
    pages: List[ParsedPage] = []
    try:
        for i, page in enumerate(doc):
            # PyMuPDF supports Arabic glyph extraction without manual transcoding.
            text = page.get_text("text")
            pages.append(ParsedPage(page_number=i + 1, text=text))
    finally:
        doc.close()

    needs_ocr = any(p.char_count < MIN_CHARS_PER_PAGE_THRESHOLD for p in pages) or not pages

    return ParseResult(
        source_path=path,
        file_type="pdf",
        pages=pages,
        extraction_method="native",
        needs_ocr=needs_ocr,
    )


def parse_docx(path: str) -> ParseResult:
    import docx  # python-docx

    document = docx.Document(path)

    parts: List[str] = []
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    # Tables often hold structured resume info (skills matrices, dates)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    full_text = "\n".join(parts)
    # DOCX has no page concept without rendering; treat as a single "page".
    page = ParsedPage(page_number=1, text=full_text)

    return ParseResult(
        source_path=path,
        file_type="docx",
        pages=[page],
        extraction_method="native",
        needs_ocr=page.char_count < MIN_CHARS_PER_PAGE_THRESHOLD,
    )


def parse_file(path: str) -> ParseResult:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return parse_pdf(path)
    if ext == ".docx":
        return parse_docx(path)
    raise UnsupportedFileTypeError(f"Unsupported file type: {ext}")
